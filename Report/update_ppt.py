import collections
import pptx
import copy

def set_shape_text_preserve_format(shape, new_text):
    if not shape.has_text_frame:
        return
    text_frame = shape.text_frame
    if len(text_frame.paragraphs) == 0:
        text_frame.text = new_text
        return

    # Split the new text into paragraphs based on \n
    lines = [line for line in new_text.split('\n') if line.strip() != '']
    if not lines:
        text_frame.clear()
        return

    # Use the first paragraph as a template for all new lines
    p_template = text_frame.paragraphs[0]._p
    parent = p_template.getparent()

    # Remove all paragraphs except the first one
    for p in text_frame.paragraphs[1:]:
        parent.remove(p._p)

    # Add duplicate paragraphs if needed
    for i in range(len(lines) - 1):
        new_p = copy.deepcopy(p_template)
        parent.append(new_p)

    # Now we have exactly len(lines) paragraphs in the shape
    # Update the text in each paragraph
    for i, line in enumerate(lines):
        p = text_frame.paragraphs[i]
        if len(p.runs) == 0:
            p.text = line
        else:
            p.runs[0].text = line
            # Clear text from other runs to avoid leftover text, while keeping the first run's formatting
            for r in p.runs[1:]:
                r.text = ""

prs = pptx.Presentation(r'p:\Projects\AIRE-TRACE\Report\Phase_0_Review(Naveen).pptx')

# Slide 1: Title
for shape in prs.slides[0].shapes:
    if shape.has_text_frame:
        if "AIRE-Trace" in shape.text:
            set_shape_text_preserve_format(shape, "AIRE-Trace: A Standardized Prospective Evaluation and Trajectory Risk Decomposition Framework for Online Agent-Failure Warning")

# Slide 2: Abstract
for shape in prs.slides[1].shapes:
    if shape.has_text_frame and "As autonomous AI agents execute" in shape.text:
        set_shape_text_preserve_format(shape, "Operating autonomous tool-augmented language agents requires early failure detection. Existing evaluations benchmark end-of-trajectory outcomes or perform retrospective root-cause attribution, conflating post-onset symptoms with genuine early warning.\nThis project establishes AIRE-Trace: a standardized, prospective, prefix-only evaluation protocol for online agent-failure monitors. We propose Trajectory Risk Decomposition (dyn_comp), which isolates dynamic execution anomalies by subtracting prior task difficulty without retrospective knowledge, yielding significant prospective performance gains under an operational 5% false-alarm budget.")

# Slide 3: Scope -> Project Progress
for shape in prs.slides[2].shapes:
    if shape.has_text_frame:
        if "Scope" in shape.text:
            set_shape_text_preserve_format(shape, "Project Progress & Work Completed")
        elif "Standardized Ingestion" in shape.text:
            set_shape_text_preserve_format(shape, "Designed AIRE-Trace protocol enforcing prefix-only limits.\nDataset Preparation: 636 retained canonical episodes.\nEstablished TRAIN/CALIB/TEST splits.\nEngineered baseline models and Trajectory Risk Decomposition (dyn_comp).\nThreshold Calibration & Evaluation under FAR <= 5%.")

# Slide 4: Objective -> Objectives
for shape in prs.slides[3].shapes:
    if shape.has_text_frame:
        if "Primary Objective:" in shape.text:
            set_shape_text_preserve_format(shape, "Establish a Prospective Evaluation Protocol (leakage-free, prefix-only).\nTrajectory Risk Decomposition: Implement dyn_comp to isolate execution-dependent anomalies from static prior task difficulty.\nOperational False-Alarm Bounding: Enforce operational thresholds governed by FAR <= 5%.\nStatistical Rigor: Validate rankings using high-resolution joint episode-cluster bootstrapping.")
        elif "Objective" in shape.text:
            set_shape_text_preserve_format(shape, "Objectives")

# Slide 5: Literature Review
for shape in prs.slides[4].shapes:
    if shape.has_text_frame:
        if "Literature Gap:" in shape.text:
            set_shape_text_preserve_format(shape, "Literature demonstrates a strict dichotomy: evaluations are either outcome-based black boxes (AgentBench) or retrospectively contaminated (TRAIL, AgentRx).\nWhile early prospective models have emerged (PrefixGuard, AgentForesight, Trajectory Guard), they lack formal onset boundaries, fail to isolate dynamic execution risk from static task difficulty, and evaluate unthresholded metrics that mask severe operational recall degradation under strict false-alarm budgets.")

# Slide 6: Problem Statement
for shape in prs.slides[5].shapes:
    if shape.has_text_frame:
        if "AI agents can fail during multi-step execution" in shape.text:
            set_shape_text_preserve_format(shape, "Autonomous language agents frequently encounter decisive divergence points during execution, eventually leading to catastrophic termination.\nStandard failure predictors map prefix features directly to failure probabilities, which confounds intrinsic task difficulty with execution-dependent failures.\nThere is currently no standardized, leakage-free framework to rigorously and prospectively evaluate early warning monitors while mathematically isolating dynamic trajectory risk under constrained operational budgets (FAR <= 5%).")

# Slide 7: Methodology
for shape in prs.slides[6].shapes:
    if shape.has_text_frame:
        if "Evaluation Workflow:" in shape.text:
            set_shape_text_preserve_format(shape, "[Prospective Prefix-Only Protocol] Disjoint testing, explicit checkpoint limits, masked post-onset events, frozen thresholds subject to FAR limit.\n[Trajectory Risk Decomposition (dyn_comp)] Framework subtracts the expected static task difficulty (the prior) from the agent's current feature-based execution risk (the posterior). This isolates the dynamic anomaly actually occurring in real time.")

# Slide 8: Results -> Results & Discussion
for shape in prs.slides[7].shapes:
    if shape.has_text_frame:
        if "Key Findings:" in shape.text:
            set_shape_text_preserve_format(shape, "Baselines suffer severe operational recall degradation (M1 Recall@FAR near zero at 5% FAR).\ndyn_comp isolates dynamic execution risk.\ndyn_comp attains a prospective Recall@FAR of 0.1167 at 0.0404 FAR.\nSignificant absolute gain of +0.1089 over the baseline M1 model.\nStatistically validated via B=10,000 episode-cluster bootstrap resamples.")

# Slide 9: Conclusion -> Summary
for shape in prs.slides[8].shapes:
    if shape.has_text_frame:
        if "Conclusion" in shape.text:
            set_shape_text_preserve_format(shape, "Summary")
        elif "Standardized evaluation is necessary" in shape.text:
            set_shape_text_preserve_format(shape, "Established AIRE-Trace: a standardized, prefix-only evaluation protocol.\nFormulated Trajectory Risk Decomposition (dyn_comp) module.\nSuccessfully isolated execution risk from static task difficulty.\nSignificant performance improvement under strict FAR budget (<= 5%).\nConfirms viability of mathematically rigorous online failure warning systems.")

# Slide 10: Publication Status -> Future Plan / Phase-II
for shape in prs.slides[9].shapes:
    if shape.has_text_frame:
        if "Publication Status" in shape.text and not "Not Yet Submitted" in shape.text:
            set_shape_text_preserve_format(shape, "Future Plan / Phase-II")
        elif "Paper: AIRE-Trace" in shape.text:
            set_shape_text_preserve_format(shape, "Implementation Refinement & Robustness Testing.\nFinal Documentation Assembly (synthesize Phase-I experiments & design).\nConference / Manuscript Submission.\nProject Demonstration (lightweight local results visualization app).\nFinal Evaluation/Report Preparation.")

prs.save(r'p:\Projects\AIRE-TRACE\Report\Phase_1_Review(Naveen).pptx')
print("Successfully saved Phase_1_Review(Naveen).pptx")
