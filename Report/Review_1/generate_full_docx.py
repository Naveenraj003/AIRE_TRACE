import docx
from docx.shared import Pt, Inches
import os

def generate_full_docx(template_path, output_path):
    doc = docx.Document(template_path)
    
    # Clear all paragraphs and tables to preserve margins and header/footer
    for p in doc.paragraphs:
        p_element = p._element
        p_element.getparent().remove(p_element)
    for tbl in doc.tables:
        tbl_element = tbl._element
        tbl_element.getparent().remove(tbl_element)
        
    def add_para(text, style='Normal', bold=False):
        p = doc.add_paragraph(style=style)
        run = p.add_run(text)
        if bold:
            run.bold = True
        return p

    add_para('ABSTRACT', style='Heading 1')
    add_para('AI agents executing long-horizon tasks require reliable monitoring to predict failures before they catastrophically conclude, but offline retrospective benchmarks do not evaluate online, mid-trajectory failure prediction. AIRE-Trace addresses this operational problem by providing a standardized prospective evaluation and trajectory risk decomposition framework. The methodology uses structured execution trajectories from datasets like tau²-Bench and TerminalBench. Since agents often fail silently, transparent evaluation rules are used to reconstruct study populations into canonical schemas and map native rewards to a universal success/failure target. Heuristic baselines, classical ML, and sequence models are considered as failure-prediction methods. The research workflow includes data ingestion, outcome mapping, strict terminal-leakage sanitization, prefix generation using outcome-independent checkpoints, feature extraction, threshold calibration, and statistical bootstrapping. The project is intentionally presented as an operational and rigorous prospective evaluation framework rather than a novel foundation model architecture, ensuring strict separation of data and active leakage sanitization to accurately benchmark agent failure-warning methods.')

    add_para('LITERATURE REVIEW', style='Heading 1')
    add_para('The research paper associated with AIRE-Trace reviews outcome-based agent benchmarks, retrospective trajectory diagnosis, and prospective failure warning systems. The cited literature shows that early agent benchmarking studies (e.g., AgentBench, GAIA) evaluate final success/failure only after execution concludes. Other studies use retrospective analysis (e.g., AgentRx, Who&When Pro, TRAIL) to localize faults post-hoc, assuming access to the complete trajectory. Recently, prospective failure-warning systems (e.g., PrefixGuard, AgentForesight, Trajectory Guard) have emerged to monitor agents online using execution prefixes. Review studies and statistical methodology literature also emphasize that validation design, episode-clustered error correction (Cameron et al.), and strict separation of calibration and test data affect the reliability of evaluation results. AIRE-Trace is positioned as a standardized evaluation protocol. It enforces strict data sanitization, prevents outcome leakage, and freezes thresholds to provide a practical, auditable framework for comparing any future online agent-monitoring models.')

    add_para('TABLE I. SUMMARY OF REVIEWED LITERATURE', style='Heading 1')
    
    table_data = [
        ["S. No.", "Authors and Year of Publication", "Approach", "Result and Accuracy"],
        ["1", "Xiao Liu et al. (2024)", "AgentBench: Evaluating LLMs as Agents.", "Focuses on outcome benchmarks evaluating final success/failure."],
        ["2", "Chang Ma et al. (2024)", "AgentBoard: Analytical Evaluation Board.", "Uses fine-grained progress-rate metrics and outcome evaluation."],
        ["3", "Gregoire Mialon et al. (2024)", "GAIA: Benchmark for General AI Assistants.", "Evaluates final answer correctness on real-world questions."],
        ["4", "Darshan Deshpande et al. (2025)", "TRAIL: Trace Reasoning and Issue Localization.", "Retrospective trace-debugging and formal taxonomy of errors."],
        ["5", "Shraddha Barke et al. (2026)", "AgentRx: Diagnosing AI Agent Failures.", "Post-hoc localization of critical failure steps using full trajectories."],
        ["6", "Jiale Liu et al. (2026)", "Who&When Pro: Failure Attribution.", "Retrospective attribution using controlled failure injection."],
        ["7", "Xinmiao Huang et al. (2026)", "PrefixGuard: Online Failure-Warning Monitors.", "Prospective early warning using LLM-agent execution traces."],
        ["8", "Boxuan Zhang et al. (2026)", "AgentForesight: Early Failure Prediction.", "Online auditing for early failure prediction in multi-agent systems."],
        ["9", "Laksh Advani (2026)", "Trajectory Guard & False-Success Detection.", "Real-time sequence anomaly detection and false-success categorization."],
        ["10", "Shengda Fan et al. (2026)", "AgentProcessBench: Diagnosing Step Quality.", "Step-level effectiveness evaluation using process annotations."],
        ["11", "Cameron et al. (2008)", "Bootstrap Inference with Clustered Errors.", "Standardized episode-level bootstrapping for confidence intervals."],
        ["12", "Sture Holm (1979)", "Sequentially Rejective Multiple Test Procedure.", "Applied for multiple-testing corrections across statistical comparisons."]
    ]
    
    table = doc.add_table(rows=1, cols=4)
    # style removed
    hdr_cells = table.rows[0].cells
    for i in range(4):
        hdr_cells[i].text = table_data[0][i]
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        
    for row_data in table_data[1:]:
        row_cells = table.add_row().cells
        for i in range(4):
            row_cells[i].text = row_data[i]
            
    add_para('The literature supports the transition from retrospective benchmarking to online prospective monitoring, but it also shows that accuracy must be interpreted with respect to strict evaluation protocols and un-leaked thresholds. AIRE-Trace therefore separates diagnostic metrics from prospective warning metrics.')

    add_para('PROBLEM STATEMENT', style='Heading 1')
    add_para('Traditional AI-agent evaluation benchmarks record the final success or failure of a task, but this does not directly indicate the agent\'s risk of failure while it is actively executing the task. Manual intervention is difficult to scale, and detecting a failure after the agent has already hallucinated or exhausted resources is too late. A practical system should analyze the agent\'s execution prefix (e.g., tool errors, retry loops, latency) mid-trajectory to produce an interpretable prospective failure warning. The major challenge is that evaluating such systems requires strict outcome-independent checkpoints and prevention of "terminal leakage" (where models peek at the end result). AIRE-Trace addresses this challenge by providing a rigorous, leakage-safe protocol and harness. The project therefore asks whether lightweight monitoring algorithms can accurately predict impending failure from execution prefixes while operating under strict operational false-alarm constraints, completely separated from retrospective advantages.')

    add_para('METHODOLOGY', style='Heading 1')
    add_para('The methodology consists of data ingestion, outcome mapping, terminal sanitization, prefix generation, feature extraction, threshold calibration, and common evaluation metrics. The implemented research model evaluates multiple algorithmic families across reconstructed study populations.')
    
    add_para('1. EXISTING ALGORITHMS', style='Heading 2')
    add_para('Heuristic Baselines (H): used as interpretable, training-free baselines; relies on deterministic rules like maximum tool errors or consecutive retries.')
    add_para('Classical ML (M): uses algorithms like Logistic Regression or XGBoost on flattened prefix features; provides a stronger, trainable baseline.')
    add_para('Sequence Models (S): uses lightweight recurrent architectures (GRU/1D-CNN) that process the ordered sequence of execution features.')

    add_para('2. RECENTLY PROPOSED ALGORITHM', style='Heading 2')
    add_para('The proposed advanced method is Trajectory Risk Decomposition (dyn_comp). This approach isolates dynamic, time-evolving execution risk from task-static baseline difficulty. A future extension could also incorporate Local LLM Monitors (L) to parse the semantic text representations of the trajectory. The current AIRE-Trace pipeline focuses on comparing classical ML (M1) and Sequence models (S1, S2) against the dyn_comp decomposition risk.')

    add_para('3. ALGORITHM USED: DYNAMIC RISK DECOMPOSITION (dyn_comp)', style='Heading 2')
    add_para('Trajectory Risk Decomposition was selected because it controls for inherent task difficulty. It calculates dynamic risk as the difference between the sequence model\'s prediction at time t and the static task risk. The calibration phase fixes the detection threshold strictly on the CALIBRATION split to limit the False Alarm Rate (FAR) to an operational budget (e.g., 5%).')

    add_para('Algorithm 1: AIRE-Trace operational failure-warning pseudocode', bold=True)
    add_para('1. Ingest raw agent trajectory and map the native reward to the canonical eventual_success flag.')
    add_para('2. Sanitize the trajectory by stripping terminal events that leak the final outcome.')
    add_para('3. Apply the outcome-independent checkpoint policy to generate the execution prefix up to time t.')
    add_para('4. Extract common, cross-dataset compatible features (e.g., latency, tool errors, retries).')
    add_para('5. Pass features to the designated monitoring model (H, M, S, or dyn_comp).')
    add_para('6. Predict the continuous risk score of eventual failure.')
    add_para('7. Apply the pre-frozen operational threshold (calibrated for FAR <= 5%).')
    add_para('8. Return the prospective failure warning alert.')

    add_para('4. METHODOLOGY POINTS', style='Heading 2')
    add_para('Final dataset: 636 canonical study episodes (TRAIN = 163, CALIBRATION = 57, TEST = 416) derived from 3,316 Tracebench rows, plus tau²-Bench for Mode B transfer.')
    add_para('Final features: Tool error count, retry loops, unique tools used, execution latency, and step sequence data.')
    add_para('Operational categories: Eventual Success and Eventual Failure.')
    add_para('Training ensures strict episode-level splitting so no prefix from a TEST episode is seen during training.')
    add_para('Thresholds are explicitly frozen during the CALIBRATION phase and never re-tuned on TEST data.')
    add_para('The target metrics include AUROC, AUPRC, and prospective Recall@FAR, evaluated with episode-clustered bootstrapping.')

    add_para('5. PSEUDOCODE OF THE ALGORITHMS', style='Heading 2')
    add_para('The following pseudocode describes the modeling families evaluated in the AIRE-Trace methodology.')

    add_para('Algorithm 1: Heuristic Baseline (H)', bold=True)
    add_para('INPUT: execution prefix up to time t, predefined heuristics')
    add_para('1. Start with the extracted trajectory prefix.')
    add_para('2. Calculate simple rule-based counts (e.g., consecutive tool errors).')
    add_para('3. If the count breaches a predefined rule threshold, predict failure.')
    add_para('4. Output prediction without requiring training.')
    add_para('OUTPUT: predicted operational success/failure risk')

    add_para('Algorithm 2: Classical ML Classification (M)', bold=True)
    add_para('INPUT: feature matrix X from canonical prefixes, labels y')
    add_para('1. Flatten the feature sequence at checkpoint t into an aggregated feature vector.')
    add_para('2. Apply episode-balanced sample weighting to ensure unbiased learning across long/short episodes.')
    add_para('3. Train a classifier (e.g., XGBoost) on the TRAIN split.')
    add_para('4. For a new execution prefix, process the features identically and apply the trained classifier.')
    add_para('OUTPUT: predicted risk and model confidence')

    add_para('Algorithm 3: Sequence Model Classification (S)', bold=True)
    add_para('INPUT: ordered feature sequence up to time t, labels y')
    add_para('1. Process the chronological sequence of events without flattening them.')
    add_para('2. Pass the sequence into a lightweight sequence architecture (GRU or 1D-CNN).')
    add_para('3. Update the hidden state iteratively to capture the temporal dependencies of the agent trajectory.')
    add_para('4. Use the final hidden state at checkpoint t to predict the prospective risk of failure.')
    add_para('OUTPUT: temporal risk prediction')

    add_para('SYSTEM DESIGN AND MODULE DEVELOPMENT', style='Heading 1')
    add_para('1. SYSTEM DESIGN', style='Heading 2')
    add_para('The AIRE-Trace system follows a strict, pipeline-driven layered design. The ingestion layer retrieves and parses raw dataset files. The mapping layer standardizes the success/failure outcomes. The sanitization layer strips terminal outcome-leakage from the events. The splitting layer enforces episode-level isolation. The prefix layer generates outcome-independent checkpoints. The feature layer builds cross-dataset compatible vectors. The model layer trains and evaluates algorithms. The evaluation layer calculates rigorous metrics (Mode A/Mode B) and performs statistical bootstrapping. This design guarantees standardized prospective evaluation without data leakage.')
    
    # Add Image
    try:
        doc.add_picture(r'p:\Projects\AIRE-TRACE\04_ieee_publication\figures\ieee\fig1_aire_trace_architecture.png', width=Inches(6.0))
        add_para('Fig. 1. System architecture of the AIRE-Trace research workflow.', bold=True)
    except Exception as e:
        print(f"Could not add image: {e}")

    add_para('2. MODULE DEVELOPMENT', style='Heading 2')
    add_para('• Ingestion module: Reads raw source files and produces Run and Event structures in the canonical schema.')
    add_para('• Outcome Mapping module: Maps native dataset outcomes to a unified eventual_success/failure boolean flag.')
    add_para('• Sanitization module: Strips terminal events that might disclose the episode outcome to monitors.')
    add_para('• Splitting module: Handles episode-level splitting (TRAIN/CALIB/TEST) ensuring disjoint datasets.')
    add_para('• Prefix Generation module: Generates prefixes using an outcome-independent checkpoint policy.')
    add_para('• Feature Extraction module: Extracts common, cross-dataset compatible features safely without leakage.')
    add_para('• Threshold Calibration module: Selects and freezes decision thresholds on calibration episodes only.')
    add_para('• Evaluation module: Computes prospective and diagnostic metrics (AUROC, AUPRC, Recall@FAR).')
    add_para('• Statistical module: Computes ranking-stability definitions (PDC, Coverage) using clustered bootstrapping.')

    add_para('REFERENCES', style='Heading 1')
    add_para('[1] X. Liu et al., "AgentBench: Evaluating LLMs as Agents," in Proc. ICLR, 2024. arXiv:2308.03688.')
    add_para('[2] C. Ma et al., "AgentBoard: An Analytical Evaluation Board of Multi-turn LLM Agents," in Proc. NeurIPS, 2024. arXiv:2401.13178.')
    add_para('[3] G. Mialon et al., "GAIA: a benchmark for General AI Assistants," in Proc. ICLR, 2024. arXiv:2311.12983.')
    add_para('[4] D. Deshpande et al., "TRAIL: Trace Reasoning and Agentic Issue Localization," arXiv preprint arXiv:2505.08638, 2025.')
    add_para('[5] S. Barke et al., "AgentRx: Diagnosing AI Agent Failures from Execution Trajectories," arXiv preprint arXiv:2602.02475, 2026.')
    add_para('[6] J. Liu et al., "Who&When Pro: Can LLMs Really Attribute Failures in AI Agents?" arXiv preprint arXiv:2607.09996, 2026.')
    add_para('[7] X. Huang et al., "PrefixGuard: From LLM-Agent Traces to Online Failure-Warning Monitors," arXiv preprint arXiv:2605.06455, 2026.')
    add_para('[8] B. Zhang et al., "AgentForesight: Online Auditing for Early Failure Prediction in Multi-Agent Systems," arXiv preprint arXiv:2605.08715, 2026.')
    add_para('[9] L. Advani, "Trajectory Guard: A Lightweight, Sequence-Aware Model for Real-Time Anomaly Detection in Agentic AI," arXiv preprint arXiv:2601.00516, 2026.')
    add_para('[10] S. Fan et al., "AgentProcessBench: Diagnosing Step-Level Process Quality in Tool-Using Agents," arXiv preprint arXiv:2603.14465, 2026.')
    add_para('[11] A. C. Cameron, J. B. Gelbach, and D. L. Miller, "Bootstrap-Based Improvements for Inference with Clustered Errors," Rev. Econ. Stat., vol. 90, no. 3, pp. 414-427, 2008.')
    add_para('[12] S. Holm, "A Simple Sequentially Rejective Multiple Test Procedure," Scand. J. Stat., vol. 6, no. 2, pp. 65-70, 1979.')

    add_para('SUMMARY', style='Heading 1')
    add_para('AIRE-Trace provides an auditable, rigorous machine-learning workflow and standardized prospective evaluation framework for online AI-agent failure warning. The project reconstructs study populations from datasets into a strict canonical schema, enforcing complete separation of training data and active leakage sanitization. By evaluating rule-based heuristics alongside classical ML and sequence-aware algorithms (like dyn_comp risk decomposition), AIRE-Trace exposes phenomena such as prospective recall collapse under operational false-alarm constraints. The methodology is deliberately built to support transparent, unbiased threshold calibration and rigorous statistical validation, providing an un-leaked environment to fairly benchmark the trajectory risk prediction capabilities of modern agents.')

    doc.save(output_path)
    print(f"Saved {output_path}")

content_template_path = r"p:\Projects\AIRE-TRACE\Report\Review_1\AirGenix_Project_Review_Content.docx"
# Overwrite the document directly as requested by the user
out_content = r"p:\Projects\AIRE-TRACE\Report\Review_1\AirGenix_Project_Review_Content.docx"
generate_full_docx(content_template_path, out_content)
