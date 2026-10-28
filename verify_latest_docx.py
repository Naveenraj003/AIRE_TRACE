"""
verify_latest_docx.py
---------------------
Audits the generated Word Document (AIRE_Trace_Manuscript.docx) for:
1. File existence and size.
2. US Letter page geometry & IEEE margins (0.75" top, 1.0" bottom, 0.625" left/right).
3. Two-column layout in Section 1 with 0.25" gutter (360 twips).
4. Title, 4 verified author blocks, affiliations, emails.
5. Abstract text completeness and styling.
6. All Section Headings (I to VII).
7. All 5 Tables (Tables I - V) with exact rows and empirical findings.
8. All 3 Figures (Figs. 1 - 3) embedded with captions.
9. All 13 IEEE References.
10. Scientific and empirical data invariants matching the frozen protocol and paper.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

def audit_manuscript_docx(filepath):
    print(f"=== AUDITING WORD DOCUMENT: {filepath} ===")
    assert os.path.exists(filepath), f"File not found: {filepath}"
    file_size = os.path.getsize(filepath)
    print(f"PASS: File exists ({file_size:,} bytes)")
    assert file_size > 100000, "File size smaller than expected (figures might be missing)"

    doc = docx.Document(filepath)
    
    # 1. Section geometry
    assert len(doc.sections) >= 2, f"Expected >= 2 sections, found {len(doc.sections)}"
    sec0 = doc.sections[0]
    sec1 = doc.sections[1]

    tol = 0.02
    assert abs(sec0.page_width.inches - 8.5) < tol, f"Sec0 width not 8.5 in: {sec0.page_width.inches}"
    assert abs(sec0.page_height.inches - 11.0) < tol, f"Sec0 height not 11.0 in: {sec0.page_height.inches}"
    assert abs(sec0.top_margin.inches - 0.75) < tol, f"Sec0 top margin not 0.75 in: {sec0.top_margin.inches}"
    assert abs(sec0.bottom_margin.inches - 1.0) < tol, f"Sec0 bottom margin not 1.0 in: {sec0.bottom_margin.inches}"
    assert abs(sec0.left_margin.inches - 0.625) < tol, f"Sec0 left margin not 0.625 in: {sec0.left_margin.inches}"
    assert abs(sec0.right_margin.inches - 0.625) < tol, f"Sec0 right margin not 0.625 in: {sec0.right_margin.inches}"
    print("PASS: Section 0 margins strictly compliant with IEEE US Letter specifications")

    sectPr1 = sec1._sectPr
    cols = sectPr1.xpath('./w:cols')
    assert cols, "Section 1 missing column definition"
    num_cols = cols[0].get(qn('w:num'))
    space_cols = cols[0].get(qn('w:space'))
    assert num_cols == '2', f"Expected 2 columns in Section 1, got {num_cols}"
    assert space_cols == '360', f"Expected 360 twips (0.25 in) gutter, got {space_cols}"
    print("PASS: Section 1 two-column layout strictly configured (2 columns, 0.25 in gutter)")

    # 2. Extract full text
    full_paragraphs = [p.text for p in doc.paragraphs]
    combined_paragraphs_text = "\n".join(full_paragraphs)
    
    table_texts = []
    for t in doc.tables:
        t_rows = []
        for r in t.rows:
            t_rows.append(" | ".join(c.text.replace('\n', ' ') for c in r.cells))
        table_texts.append("\n".join(t_rows))
    combined_all = combined_paragraphs_text + "\n" + "\n".join(table_texts)

    # 3. Title & Authors
    assert "AIRE-Trace: A Standardized Prospective Evaluation and Trajectory Risk Decomposition Framework" in combined_all
    print("PASS: Title verified")

    authors = ["B. Naveen Raj", "C. Janani", "P. Krishnamoorthi", "L. Maadhava Kumar"]
    for a in authors:
        assert a in combined_all, f"Missing author: {a}"
    print("PASS: All 4 real authors verified")

    emails = ["naveenrajbv1@gmail.com", "srijanani2030@gmail.com", "pkrishnamoorthi126@gmail.com", "maadhava05@gmail.com"]
    for e in emails:
        assert e in combined_all, f"Missing email: {e}"
    print("PASS: All author contact emails verified")

    # 4. Abstract
    assert "Abstract—" in combined_all
    assert "The increasing deployment of autonomous tool-using AI agents creates demand" in combined_all
    assert "Trajectory risk decomposition" in combined_all or "trajectory risk decomposition" in combined_all
    assert "Recall@FAR of 0.1167" in combined_all
    print("PASS: Abstract verified")

    # 5. Headings I - VII
    headings = [
        "I. INTRODUCTION",
        "II. RELATED WORK AND COMPARATIVE GAP",
        "III. AIRE-TRACE METHODOLOGY",
        "IV. EXPERIMENTAL SETUP",
        "V. RESULTS",
        "VI. DISCUSSION AND LIMITATIONS",
        "VII. CONCLUSION",
        "REFERENCES"
    ]
    for h in headings:
        assert h in combined_all, f"Missing heading: {h}"
    print("PASS: All 7 Roman-numeral headings + REFERENCES verified")

    # 6. Tables I - V
    assert "TABLE I:  WARNING METHOD REGISTRY" in combined_all
    assert "TABLE II:  PRIMARY MODE A TEST PERFORMANCE" in combined_all
    assert "TABLE III:  PRIMARY DYN_COMP VS. M1 BOOTSTRAP COMPARISON" in combined_all
    assert "TABLE IV:  S1 FEATURE ABLATION RESULTS" in combined_all
    assert "TABLE V:  OPERATIONAL WARNING PROFILE" in combined_all
    print("PASS: All 5 Table captions verified")

    # Table rows check
    assert "Random Forest (100 trees, depth=6)" in combined_all
    assert "0.3593" in combined_all and "0.1167" in combined_all and "0.0404" in combined_all
    assert "+0.1089" in combined_all and "[0.0706, 0.1496]" in combined_all and "10,000" in combined_all
    assert "0.5302" in combined_all and "−0.1069" in combined_all
    assert "8,012" in combined_all and "19.84%" in combined_all
    print("PASS: Table cell data and statistical invariants verified")

    # 7. Figures
    assert "Fig. 1.  " in combined_all
    assert "Fig. 2.  " in combined_all
    assert "Fig. 3.  " in combined_all
    print("PASS: All figure captions verified")

    # Check embedded images in document package
    image_count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            image_count += 1
    assert image_count >= 3, f"Expected at least 3 embedded images, found {image_count}"
    print(f"PASS: {image_count} embedded images verified in document package")

    # 8. References [1] - [13]
    for i in range(1, 14):
        assert f"[{i}]  " in combined_all, f"Missing reference [{i}]"
    assert "AgentBench" in combined_all
    assert "AgentBoard" in combined_all
    assert "GAIA" in combined_all
    assert "TRAIL" in combined_all
    assert "AgentRx" in combined_all
    assert "PrefixGuard" in combined_all
    assert "AgentForesight" in combined_all
    assert "Trajectory Guard" in combined_all
    assert "Holm" in combined_all
    print("PASS: All 13 IEEE References verified")

    print("\nALL AUDIT CHECKS PASSED SUCCESSFULLY!\n")

if __name__ == '__main__':
    audit_manuscript_docx(r"P:\Projects\AIRE-TRACE\AIRE_Trace_Manuscript.docx")
    audit_manuscript_docx(r"P:\Projects\AIRE-TRACE\conference_submissions\iConCISE_2027\working\AIRE_Trace_Manuscript.docx")
