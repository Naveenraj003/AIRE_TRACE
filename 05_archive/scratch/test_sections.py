import docx
from docx.shared import Inches, Pt
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cols(section, n_cols=2, space=360):
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        col = cols[0]
    else:
        col = OxmlElement('w:cols')
        sectPr.append(col)
    col.set(qn('w:num'), str(n_cols))
    col.set(qn('w:space'), str(space))

doc = docx.Document()
# Section 0: 1-col
s0 = doc.sections[0]
s0.page_width = Inches(8.5)
s0.page_height = Inches(11.0)
s0.top_margin = Inches(0.75)
s0.bottom_margin = Inches(1.0)
s0.left_margin = Inches(0.625)
s0.right_margin = Inches(0.625)
doc.add_paragraph("Title across 1 column")

# Section 1: 2-col
s1 = doc.add_section(WD_SECTION_START.CONTINUOUS)
s1.top_margin = Inches(0.75)
s1.bottom_margin = Inches(1.0)
s1.left_margin = Inches(0.625)
s1.right_margin = Inches(0.625)
set_cols(s1, 2, 360)
doc.add_paragraph("Abstract and body column 1 and 2 text...")

# Section 2: 1-col for Wide Table
s2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
s2.top_margin = Inches(0.75)
s2.bottom_margin = Inches(1.0)
s2.left_margin = Inches(0.625)
s2.right_margin = Inches(0.625)
set_cols(s2, 1, 0)
doc.add_paragraph("Wide Table III across both columns")

# Section 3: 2-col to resume body
s3 = doc.add_section(WD_SECTION_START.CONTINUOUS)
s3.top_margin = Inches(0.75)
s3.bottom_margin = Inches(1.0)
s3.left_margin = Inches(0.625)
s3.right_margin = Inches(0.625)
set_cols(s3, 2, 360)
doc.add_paragraph("Body resumes in 2 columns...")

doc.save("scratch/test_multi_section.docx")
print("Multi-section continuous breaks passed!")
